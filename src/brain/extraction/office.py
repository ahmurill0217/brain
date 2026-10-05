"""Word and PowerPoint files to markdown.

Both formats go through the library that actually reads them: mammoth for
.docx (to HTML, then markdownify to markdown) and python-pptx for .pptx. The
output follows what markitdown produced when brain went through it, so swapping
it out did not move chunk boundaries or retrieval:

  .docx  headings as `#`, emphasis, lists, tables, and links as markdown.
         Images are left out: they become their own sections, summarized by
         the vision model, and a placeholder here would only be indexed noise.
  .pptx  a `<!-- Slide number: N -->` marker per slide, the title as `#`, text
         and tables in reading order (top to bottom, then left to right, groups
         included), picture alt text, and speaker notes under `### Notes:`.
         A chart becomes "[chart omitted]": rendering one's series is slow on
         a chart-heavy deck, and its numbers are rarely what anyone searches.

Images themselves are pulled from the zip in `extract.py`, not here.
"""

from __future__ import annotations

import re
from typing import IO, Any

CHART_PLACEHOLDER = "\n\n[chart omitted]\n\n"


def docx_to_markdown(file: IO[bytes]) -> str:
    """Raises on a file that is not a readable .docx; the caller falls back."""
    import mammoth
    from markdownify import ATX, markdownify

    html = mammoth.convert_to_html(file).value
    return _normalize(markdownify(html, heading_style=ATX, strip=["img"]))


def pptx_to_markdown(file: IO[bytes]) -> str:
    """Raises on a file that is not a readable .pptx; the caller decides."""
    import pptx

    presentation = pptx.Presentation(file)
    markdown = ""
    for number, slide in enumerate(presentation.slides, 1):
        markdown += f"\n\n<!-- Slide number: {number} -->\n"
        title = slide.shapes.title
        for shape in _reading_order(slide.shapes):
            markdown += _shape_markdown(shape, title)
        markdown = markdown.strip()

        # PowerPoint attaches a notes slide as soon as the notes pane has been
        # opened, so having one says nothing about there being notes to read.
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = (frame.text or "") if frame is not None else ""
            if notes.strip():
                markdown = (markdown + "\n\n### Notes:\n" + notes).strip()
    return _normalize(markdown)


def _normalize(markdown: str) -> str:
    """No trailing spaces, and at most one blank line in a row."""
    lines = "\n".join(line.rstrip() for line in re.split(r"\r?\n", markdown))
    return re.sub(r"\n{3,}", "\n\n", lines).strip()


def _reading_order(shapes: Any) -> list[Any]:
    """Top to bottom, then left to right. Unpositioned shapes come first."""
    return sorted(
        shapes,
        key=lambda s: (
            float("-inf") if s.top is None else s.top,
            float("-inf") if s.left is None else s.left,
        ),
    )


def _shape_markdown(shape: Any, title: Any) -> str:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    parts: list[str] = []
    if _is_picture(shape):
        parts.append(_picture_markdown(shape))
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        parts.append(_table_markdown(shape.table))
    if shape.has_chart:
        parts.append(CHART_PLACEHOLDER)
    elif shape.has_text_frame:
        text = shape.text or ""
        if shape == title:
            if text.strip():
                parts.append("# " + text.lstrip() + "\n")
        else:
            parts.append(text + "\n")
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        parts.extend(_shape_markdown(child, title) for child in _reading_order(shape.shapes))
    return "".join(parts)


def _is_picture(shape: Any) -> bool:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return True
    return shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and hasattr(shape, "image")


def _picture_markdown(shape: Any) -> str:
    """The picture's alt text, falling back to its shape name."""
    try:
        alt_text = shape._element._nvXxPr.cNvPr.attrib.get("descr", "")
    except AttributeError:
        alt_text = ""
    alt_text = re.sub(r"[\r\n\[\]]", " ", alt_text or shape.name)
    alt_text = re.sub(r"\s+", " ", alt_text).strip()
    file_name = re.sub(r"\W", "", shape.name) + ".jpg"
    return f"\n![{alt_text}]({file_name})\n"


def _table_markdown(table: Any) -> str:
    """The first row is the header, as PowerPoint tables are laid out."""
    rows = [[_cell_text(cell) for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    lines = [
        "| " + " | ".join(rows[0]) + " |",
        "| " + " | ".join("---" for _ in rows[0]) + " |",
        *("| " + " | ".join(row) + " |" for row in rows[1:]),
    ]
    return "\n".join(lines) + "\n"


def _cell_text(cell: Any) -> str:
    return re.sub(r"\s+", " ", cell.text or "").strip().replace("|", "\\|")
