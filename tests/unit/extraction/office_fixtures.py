"""Word and PowerPoint files built in memory, covering what extraction must read.

Shared by the docx and pptx tests. Each builder returns a BytesIO positioned at
zero.
"""

from __future__ import annotations

import io
import struct
import zlib

from docx import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def png_1x1() -> bytes:
    """A minimal valid 1x1 white PNG."""
    signature = b"\x89PNG\r\n\x1a\n"

    def _chunk(ctype: bytes, data: bytes) -> bytes:
        crc = zlib.crc32(ctype + data) & 0xFFFFFFFF
        return struct.pack(">I", len(data)) + ctype + data + struct.pack(">I", crc)

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00\xff\xff\xff")
    return signature + _chunk(b"IHDR", ihdr) + _chunk(b"IDAT", idat) + _chunk(b"IEND", b"")


def _save(document: object) -> io.BytesIO:
    buf = io.BytesIO()
    document.save(buf)  # type: ignore[attr-defined]
    buf.seek(0)
    return buf


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    link.append(run)
    paragraph._p.append(link)


def rich_docx() -> io.BytesIO:
    """Headings, emphasis, both list kinds, a table, a link, and an image."""
    doc = WordDocument()
    doc.add_heading("Benefits Overview", level=1)
    intro = doc.add_paragraph("The plan year starts on ")
    intro.add_run("July 1").bold = True
    intro.add_run(" and enrollment is ")
    intro.add_run("required").italic = True
    intro.add_run(".")

    doc.add_heading("Eligibility", level=2)
    doc.add_paragraph("Full-time employees", style="List Bullet")
    doc.add_paragraph("Employees with proof of other coverage", style="List Bullet")
    doc.add_paragraph("Submit the form", style="List Number")
    doc.add_paragraph("Wait for confirmation", style="List Number")

    table = doc.add_table(rows=3, cols=2)
    for row, (tier, amount) in enumerate(
        [("Tier", "Payment"), ("Employee", "$41.67"), ("Family", "$83.34")]
    ):
        table.cell(row, 0).text = tier
        table.cell(row, 1).text = amount

    contact = doc.add_paragraph("Questions? See ")
    _add_hyperlink(contact, "https://hr.example/benefits", "the benefits portal")
    # The curly apostrophe is deliberate: non-ASCII must survive conversion.
    contact.add_run(" or email the People team — they’ll help.")  # noqa: RUF001

    doc.add_picture(io.BytesIO(png_1x1()))
    return _save(doc)


def rich_pptx() -> io.BytesIO:
    """A titled slide, a table, a group, notes, a picture with alt text, and a chart."""
    prs = Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Review"
    slide.placeholders[1].text = "Revenue grew\nCosts held flat"
    slide.notes_slide.notes_text_frame.text = "Mention the hiring plan."

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Headcount"
    rows = [("Team", "People"), ("Research", "12"), ("Platform", "7")]
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    first = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    first.text_frame.text = "Grouped first"
    second = group.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
    second.text_frame.text = "Grouped second"
    picture = slide.shapes.add_picture(io.BytesIO(png_1x1()), Inches(5), Inches(1))
    picture._element._nvXxPr.cNvPr.set("descr", "Org chart [draft]")
    # Opened but empty notes pane: must not produce a Notes heading.
    _ = slide.notes_slide

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Revenue chart"
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3"]
    chart_data.add_series("Revenue", (100, 200, 300))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), chart_data
    )
    return _save(prs)
