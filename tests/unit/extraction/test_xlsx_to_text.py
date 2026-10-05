from __future__ import annotations

import io
from typing import cast
from unittest.mock import patch

import openpyxl
import pytest
from openpyxl.worksheet.worksheet import Worksheet

from brain.config import BrainSettings
from brain.extraction.extract import (
    _sheet_to_csv,
    xlsx_sheet_extraction,
    xlsx_sheets_to_csv,
    xlsx_to_text,
)

_NO_CELL_LIMIT = 10_000_000


def make_xlsx(sheets: dict[str, list[list[str]]]) -> io.BytesIO:
    """An in-memory xlsx from a dict of sheet_name -> matrix of strings."""
    wb = openpyxl.Workbook()
    if wb.active is not None:
        wb.remove(cast(Worksheet, wb.active))
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class TestXlsxToText:
    def test_formula_strings_are_not_indexed(self, extraction_settings: BrainSettings) -> None:
        # Workbooks are read with data_only=True, so a cell yields its cached
        # calculated value rather than formula source. A formula no spreadsheet
        # app ever evaluated has no cached value and is omitted: the accepted
        # tradeoff for not indexing raw "=SUM(...)" strings as prose.
        wb = openpyxl.Workbook()
        ws = cast(Worksheet, wb.active)
        ws.append(["Total", "=SUM(B2:B10)"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        result = xlsx_to_text(buf, settings=extraction_settings)
        assert "=SUM" not in result
        assert "Total" in result

    def test_single_sheet_basic(self, extraction_settings: BrainSettings) -> None:
        xlsx = make_xlsx({"Sheet1": [["Name", "Age"], ["Alice", "30"], ["Bob", "25"]]})
        result = xlsx_to_text(xlsx, settings=extraction_settings)
        lines = [line for line in result.strip().split("\n") if line.strip()]
        assert len(lines) == 3
        assert "Name" in lines[0]
        assert "Alice" in lines[1]
        assert "Bob" in lines[2]

    def test_multiple_sheets_separated(self, extraction_settings: BrainSettings) -> None:
        xlsx = make_xlsx({"Sheet1": [["a", "b"]], "Sheet2": [["c", "d"]]})
        result = xlsx_to_text(xlsx, settings=extraction_settings)
        assert "\n\n" in result
        parts = result.split("\n\n")
        assert any("a" in p for p in parts)
        assert any("c" in p for p in parts)

    def test_empty_cells(self, extraction_settings: BrainSettings) -> None:
        xlsx = make_xlsx({"Sheet1": [["a", "", "b"], ["", "c", ""]]})
        result = xlsx_to_text(xlsx, settings=extraction_settings)
        lines = [line for line in result.strip().split("\n") if line.strip()]
        assert len(lines) == 2

    def test_commas_in_cells_are_quoted(self, extraction_settings: BrainSettings) -> None:
        xlsx = make_xlsx({"Sheet1": [["hello, world", "normal"]]})
        assert '"hello, world"' in xlsx_to_text(xlsx, settings=extraction_settings)

    def test_empty_workbook(self, extraction_settings: BrainSettings) -> None:
        assert xlsx_to_text(make_xlsx({"Sheet1": []}), settings=extraction_settings).strip() == ""

    def test_long_empty_row_run_capped(self, extraction_settings: BrainSettings) -> None:
        xlsx = make_xlsx({"Sheet1": [["header"], [""], [""], [""], [""], ["data"]]})
        result = xlsx_to_text(xlsx, settings=extraction_settings)
        lines = [line for line in result.strip().split("\n") if line.strip()]
        # 4 empty rows capped to 2: header + 2 empty + data = 4 lines
        assert len(lines) == 4
        assert "header" in lines[0]
        assert "data" in lines[-1]

    def test_long_empty_col_run_capped(self, extraction_settings: BrainSettings) -> None:
        xlsx = make_xlsx({"Sheet1": [["a", "", "", "", "b"], ["c", "", "", "", "d"]]})
        result = xlsx_to_text(xlsx, settings=extraction_settings)
        lines = [line for line in result.strip().split("\n") if line.strip()]
        assert len(lines) == 2
        # a + 2 empty + b is 4 fields, so 3 commas, not 4.
        assert lines[0].strip().count(",") == 3

    def test_short_empty_runs_kept(self, extraction_settings: BrainSettings) -> None:
        xlsx = make_xlsx({"Sheet1": [["a", "b"], ["", ""], ["", ""], ["c", "d"]]})
        result = xlsx_to_text(xlsx, settings=extraction_settings)
        lines = [line for line in result.strip().split("\n") if line.strip()]
        assert len(lines) == 4

    def test_bad_zip_file_returns_empty(self, extraction_settings: BrainSettings) -> None:
        bad_file = io.BytesIO(b"not a zip file")
        assert xlsx_to_text(bad_file, settings=extraction_settings, file_name="test.xlsx") == ""

    def test_bad_zip_tilde_file_returns_empty(self, extraction_settings: BrainSettings) -> None:
        bad_file = io.BytesIO(b"not a zip file")
        assert xlsx_to_text(bad_file, settings=extraction_settings, file_name="~$temp.xlsx") == ""

    def test_large_sparse_sheet(self, extraction_settings: BrainSettings) -> None:
        rows: list[list[str]] = [["row1_data"]]
        rows.extend([[""] for _ in range(10)])
        rows.append(["row2_data"])
        result = xlsx_to_text(make_xlsx({"Sheet1": rows}), settings=extraction_settings)
        lines = [line for line in result.strip().split("\n") if line.strip()]
        # 10 empty rows capped to 2: row1_data + 2 empty + row2_data = 4
        assert len(lines) == 4

    def test_quotes_in_cells(self, extraction_settings: BrainSettings) -> None:
        xlsx = make_xlsx({"Sheet1": [['say "hello"', "normal"]]})
        # csv.writer escapes quotes by doubling them.
        assert '""hello""' in xlsx_to_text(xlsx, settings=extraction_settings)

    def test_each_row_is_separate_line(self, extraction_settings: BrainSettings) -> None:
        xlsx = make_xlsx(
            {"Sheet1": [["r1c1", "r1c2"], ["r2c1", "r2c2"], ["r3c1", "r3c2"]]}
        )
        result = xlsx_to_text(xlsx, settings=extraction_settings)
        lines = [line for line in result.strip().split("\n") if line.strip()]
        assert len(lines) == 3
        assert "r1c1" in lines[0] and "r1c2" in lines[0]
        assert "r3c1" in lines[2]


class TestSheetToCsvJaggedRows:
    """openpyxl's read-only mode yields rows of differing widths when trailing
    cells are empty. These go through `_sheet_to_csv` directly because
    `make_xlsx` normalizes row widths via `ws.append`."""

    def test_shorter_trailing_rows_padded_in_output(self) -> None:
        csv_text = _sheet_to_csv(
            iter([("A", "B", "C"), ("X", "Y"), ("P",)]), max_cells=_NO_CELL_LIMIT
        )
        assert csv_text.split("\n") == ["A,B,C", "X,Y,", "P,,"]

    def test_shorter_leading_row_padded_in_output(self) -> None:
        csv_text = _sheet_to_csv(iter([("A",), ("X", "Y", "Z")]), max_cells=_NO_CELL_LIMIT)
        assert csv_text.split("\n") == ["A,,", "X,Y,Z"]

    def test_no_index_error_on_jagged_rows(self) -> None:
        """Regression: the dense-matrix version raised IndexError when a later
        row was shorter than an earlier one whose extra columns were empty."""
        csv_text = _sheet_to_csv(
            iter([("A", "", "", "B"), ("X", "Y")]), max_cells=_NO_CELL_LIMIT
        )
        assert csv_text.split("\n") == ["A,,,B", "X,Y,,"]


class TestSheetToCsvStreaming:
    """The memory-safe streaming contract: empty rows are skipped cheaply,
    empty row/column runs collapse to at most 2, and a sheet with no data
    returns the empty string."""

    def test_empty_rows_between_data_capped_at_two(self) -> None:
        csv_text = _sheet_to_csv(
            iter([("A", "B"), *[(None, None)] * 5, ("C", "D")]), max_cells=_NO_CELL_LIMIT
        )
        assert csv_text.split("\n") == ["A,B", ",", ",", "C,D"]

    def test_empty_rows_at_or_below_cap_preserved(self) -> None:
        csv_text = _sheet_to_csv(
            iter([("A", "B"), (None, None), (None, None), ("C", "D")]),
            max_cells=_NO_CELL_LIMIT,
        )
        assert csv_text.split("\n") == ["A,B", ",", ",", "C,D"]

    def test_empty_column_run_capped_at_two(self) -> None:
        csv_text = _sheet_to_csv(
            iter([("A", None, None, None, None, "B"), ("C", None, None, None, None, "D")]),
            max_cells=_NO_CELL_LIMIT,
        )
        assert csv_text.split("\n") == ["A,,,B", "C,,,D"]

    def test_completely_empty_stream_returns_empty_string(self) -> None:
        assert _sheet_to_csv(iter([]), max_cells=_NO_CELL_LIMIT) == ""

    def test_all_rows_empty_returns_empty_string(self) -> None:
        csv_text = _sheet_to_csv(
            iter([(None, None), ("", ""), (None,)]), max_cells=_NO_CELL_LIMIT
        )
        assert csv_text == ""

    def test_trailing_empty_rows_dropped(self) -> None:
        csv_text = _sheet_to_csv(
            iter([("A",), ("B",), (None,), (None,), (None,)]), max_cells=_NO_CELL_LIMIT
        )
        assert csv_text.split("\n") == ["A", "B"]

    def test_leading_empty_rows_capped_at_two(self) -> None:
        csv_text = _sheet_to_csv(
            iter([*[(None, None)] * 5, ("A", "B")]), max_cells=_NO_CELL_LIMIT
        )
        assert csv_text.split("\n") == [",", ",", "A,B"]

    def test_cell_cap_truncates_and_appends_marker(self) -> None:
        """Past the cap, scanning stops and a marker row is appended so the
        cut-off is visible downstream instead of silent."""
        csv_text = _sheet_to_csv(
            iter([("A", "B", "C"), ("D", "E", "F"), ("G", "H", "I"), ("J", "K", "L")]),
            max_cells=5,
        )
        lines = csv_text.split("\n")
        assert lines[-1] == "[truncated: sheet exceeded cell limit]"
        # The first two rows (6 cells) trip the cap after row 2; rows 3 and 4
        # are never scanned.
        assert "G" not in csv_text
        assert "J" not in csv_text

    def test_cell_cap_not_hit_no_marker(self) -> None:
        csv_text = _sheet_to_csv(iter([("A", "B"), ("C", "D")]), max_cells=_NO_CELL_LIMIT)
        assert "[truncated" not in csv_text

    def test_cell_cap_comes_from_settings(self, extraction_settings: BrainSettings) -> None:
        capped = extraction_settings.model_copy(update={"max_xlsx_cells_per_sheet": 2})
        xlsx = make_xlsx({"S1": [["a", "b"], ["c", "d"], ["e", "f"]]})
        assert "[truncated" in xlsx_to_text(xlsx, settings=capped)


class TestXlsxSheetExtraction:
    def test_one_tuple_per_sheet(self, extraction_settings: BrainSettings) -> None:
        xlsx = make_xlsx(
            {
                "Revenue": [["Month", "Amount"], ["Jan", "100"]],
                "Expenses": [["Category", "Cost"], ["Rent", "500"]],
            }
        )
        sheets = xlsx_sheet_extraction(xlsx, settings=extraction_settings)
        assert [title for _csv, title in sheets] == ["Revenue", "Expenses"]
        assert "Jan" in sheets[0][0]
        assert "Rent" in sheets[1][0]

    def test_tuple_structure_is_csv_text_then_title(
        self, extraction_settings: BrainSettings
    ) -> None:
        """Pin the tuple order so callers that unpack positionally do not
        silently break."""
        sheets = xlsx_sheet_extraction(
            make_xlsx({"MySheet": [["a", "b"]]}), settings=extraction_settings
        )
        csv_text, title = sheets[0]
        assert title == "MySheet"
        assert "a" in csv_text

    def test_empty_sheet_included_with_empty_csv(
        self, extraction_settings: BrainSettings
    ) -> None:
        sheets = xlsx_sheet_extraction(
            make_xlsx({"Data": [["a", "b"]], "Empty": []}), settings=extraction_settings
        )
        assert [title for _csv, title in sheets] == ["Data", "Empty"]
        assert next(c for c, title in sheets if title == "Empty") == ""

    def test_empty_workbook_returns_one_tuple_per_sheet(
        self, extraction_settings: BrainSettings
    ) -> None:
        sheets = xlsx_sheet_extraction(
            make_xlsx({"Sheet1": [], "Sheet2": []}), settings=extraction_settings
        )
        assert sheets == [("", "Sheet1"), ("", "Sheet2")]

    def test_bad_zip_returns_empty_list(self, extraction_settings: BrainSettings) -> None:
        bad_file = io.BytesIO(b"not a zip file")
        assert xlsx_sheet_extraction(bad_file, settings=extraction_settings, file_name="t.xlsx") == []

    def test_bad_zip_tilde_file_returns_empty_list(
        self, extraction_settings: BrainSettings
    ) -> None:
        """`~$`-prefixed files are Excel lock files: never valid workbooks, and
        not worth a warning."""
        bad_file = io.BytesIO(b"not a zip file")
        assert (
            xlsx_sheet_extraction(bad_file, settings=extraction_settings, file_name="~$t.xlsx")
            == []
        )

    def test_known_openpyxl_bug_returns_empty(self, extraction_settings: BrainSettings) -> None:
        """openpyxl rejects font family values > 14 with 'Max value is 14'.
        Skip the file rather than fail the whole batch."""
        with patch(
            "brain.extraction.extract.openpyxl.load_workbook",
            side_effect=ValueError("Max value is 14"),
        ):
            sheets = xlsx_sheet_extraction(
                io.BytesIO(b""), settings=extraction_settings, file_name="bad_font.xlsx"
            )
        assert sheets == []

    def test_unknown_openpyxl_error_propagates(
        self, extraction_settings: BrainSettings
    ) -> None:
        """Only the known-bug list is swallowed; anything else is a real failure."""
        with (
            patch(
                "brain.extraction.extract.openpyxl.load_workbook",
                side_effect=ValueError("something genuinely new"),
            ),
            pytest.raises(ValueError, match="genuinely new"),
        ):
            xlsx_sheet_extraction(
                io.BytesIO(b""), settings=extraction_settings, file_name="x.xlsx"
            )

    def test_csv_content_matches_xlsx_to_text_per_sheet(
        self, extraction_settings: BrainSettings
    ) -> None:
        data = [["Name", "Age"], ["Alice", "30"]]
        expected = xlsx_to_text(make_xlsx({"People": data}), settings=extraction_settings)
        sheets = xlsx_sheet_extraction(
            make_xlsx({"People": data}), settings=extraction_settings
        )
        assert sheets[0][0].strip() == expected.strip()

    def test_sheet_title_with_special_chars_preserved(
        self, extraction_settings: BrainSettings
    ) -> None:
        """Spaces, punctuation, and unicode survive verbatim: the title becomes
        a link anchor downstream."""
        xlsx = make_xlsx({"Q1 Revenue (USD)": [["a", "b"]], "Données": [["c", "d"]]})
        titles = [title for _csv, title in xlsx_sheet_extraction(xlsx, settings=extraction_settings)]
        assert titles == ["Q1 Revenue (USD)", "Données"]


class TestXlsxSheetsToCsv:
    """The faithful counterpart: (title, csv_text), empty rows dropped, nothing
    collapsed, nothing trimmed. This is what feeds TabularSection."""

    def test_tuple_order_is_title_then_csv(self) -> None:
        sheets = xlsx_sheets_to_csv(make_xlsx({"MySheet": [["a", "b"]]}))
        title, csv_text = sheets[0]
        assert title == "MySheet"
        assert csv_text == "a,b\n"

    def test_one_entry_per_non_empty_sheet(self) -> None:
        sheets = xlsx_sheets_to_csv(
            make_xlsx({"Data": [["a"]], "Empty": [], "More": [["b"]]})
        )
        assert [title for title, _ in sheets] == ["Data", "More"]

    def test_columns_are_not_trimmed(self) -> None:
        """Unlike the text path, wide empty column runs are preserved: the
        chunker reads this as a table, where column position is meaning."""
        sheets = xlsx_sheets_to_csv(make_xlsx({"S1": [["a", "", "", "", "b"]]}))
        assert sheets[0][1] == "a,,,,b\n"

    def test_empty_rows_dropped(self) -> None:
        sheets = xlsx_sheets_to_csv(make_xlsx({"S1": [["a"], [""], [""], ["b"]]}))
        assert sheets[0][1] == "a\nb\n"

    def test_bad_zip_returns_empty_list(self) -> None:
        assert xlsx_sheets_to_csv(io.BytesIO(b"not a zip"), file_name="t.xlsx") == []
